from __future__ import annotations

import argparse
import math
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import openpyxl
import pandas as pd


REPO_ROOT = Path(__file__).resolve().parents[1]
REAL_CASE_DIR = REPO_ROOT / "data" / "Instance" / "real_case_2026-05-21"
RAW_DIMENSIONS = REPO_ROOT / "data" / "raw data" / "尺寸整理.xlsx"
SAMPLING_WORKBOOK = REAL_CASE_DIR / "装车编号算例抽样.xlsx"
OPTIONAL_QUANTITY = 50


@dataclass(frozen=True)
class Dimension:
    length: int
    height: int
    source: str


@dataclass(frozen=True)
class RawRow:
    row_number: int
    project: str
    common_name: str
    model: str
    length: int
    height: int
    norm_common: str
    norm_model: str
    code_model: str


def normalize(value: object) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    text = unicodedata.normalize("NFKC", str(value)).strip().lower()
    return re.sub(r"\s+", "", text)


def clean_code(value: object) -> str:
    text = normalize(value)
    return re.sub(r"^(lzw|sc|cc|jkc)", "", text)


def parse_single_int(value: object) -> int | None:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if isinstance(value, (int, float)):
        return int(round(float(value)))
    matches = re.findall(r"\d+(?:\.\d+)?", str(value))
    if len(matches) != 1:
        return None
    return int(round(float(matches[0])))


def project_matches(raw_project: str, program: str) -> bool:
    aliases = {
        "上汽通用五菱": ["五菱"],
        "重庆长安": ["长安系列", "河北长安", "定州长安"],
        "永川长城": ["长城"],
        "长安福特": ["福特"],
        "潍柴汽车": ["潍柴"],
        "东风小康": ["东风渝安", "重庆渝安"],
    }
    keys = aliases.get(str(program).strip(), [str(program).strip()])
    raw_project = str(raw_project).strip()
    return any(key in raw_project for key in keys)


def restored_vehicle_type(program: str, model: str, detail_model: str = "") -> tuple[str, str]:
    """Recover the original common model name from the detailed vehicle model code."""
    detail_key = normalize(detail_model)
    if program == "上汽通用五菱":
        by_detail = {
            normalize("lzw6389"): "五菱之光",
            normalize("LZW5022"): "五菱之光厢货",
            normalize("LZW1020"): "五菱之光货车",
            normalize("LZW5021"): "宏光V",
            normalize("LZW6448"): "宏光V",
            normalize("LZW6443"): "新宏光",
            normalize("LZW6449"): "新款宏光S",
        }
        return program, by_detail.get(detail_key, model)

    if program == "重庆长安":
        by_detail = {
            normalize("SC7155"): "CS55",
            normalize("SC7155ACB6.CNH2002.WC2"): "CS55",
            normalize("SC7155ADA6"): "CS55蓝鲸版",
            normalize("SC7155ADB6"): "CS55蓝鲸版",
            normalize("SC7153"): "X5",
            normalize("sc7163"): "X5",
            normalize("SC6493"): "sc6493",
            normalize("SC7151"): "引力",
            normalize("SC7164G"): "新CS35PLUS",
            normalize("SC7144G"): "新CS35PLUS蓝鲸版",
            normalize("SC6482"): "新欧尚",
            normalize("SC7169K"): "新逸动PLUS",
            normalize("6478"): "欧尚X7",
            normalize("SC7001A"): "电动奔奔",
            normalize("SC7157"): "睿骋CC",
            normalize("SC7164H"): "科赛5",
            normalize("SC6491CBA6.CNA4001.WC8"): "科赛PLUS",
            normalize("SC7145K3"): "第二代逸动PLUS",
            normalize("SC7145"): "逸动PLUS蓝鲸版",
            normalize("SC7157ADB6"): "锐骋CC蓝鲸版",
            normalize("SC7157ADD6.CNH4001.WC2"): "锐骋CC蓝鲸版",
        }
        return program, by_detail.get(detail_key, model)

    if program == "永川长城":
        by_detail = {
            normalize("CC1030QA20A"): "CC1030QA20A",
            normalize("CC1030QA40A"): "CC1030QA40A",
            normalize("CC1030QA60A"): "CC1030QA60A",
            normalize("CC1030QS00B"): "CC1030QS00B",
            normalize("CC1032QS60A"): "CC1032QS60A-",
            normalize("CC1030QD51C"): "出口车大单排",
            normalize("CC1030QS52C-"): "出口车小双",
            normalize("CC1030QS72C-"): "出口车小双",
            normalize("CC1033QS62C-"): "出口车小双",
            normalize("CC2030BE21B"): "坦克300",
            normalize("CC2030BE21B-"): "坦克300",
            normalize("CC1030QA60B"): "长城炮商用大双",
            normalize("CC1030QA40B"): "长城炮大双",
            normalize("CC1030QS60B"): "长城炮小双",
            normalize("CC1032QS20B"): "长城炮越野皮卡",
            normalize("CC1032QS60B"): "长城炮越野皮卡",
            normalize("CC1030QS20B"): "长城皮卡",
        }
        return program, by_detail.get(detail_key, model)

    if program == "东风小康":
        by_detail = {
            normalize("B-风光500"): "B-风光500",
            normalize("B-风光580红星版"): "B-风光580红星版",
            normalize("C-C31"): "C-C31",
            normalize("C-C32"): "C-C32",
            normalize("C-C51L"): "C-C51L",
            normalize("C-C52L"): "C-C52L",
            normalize("C-D01"): "C-D01",
            normalize("C-D02"): "C-D02",
        }
        return program, by_detail.get(detail_key, model)

    if program == "长安福特":
        by_detail = {
            normalize("88AJK"): "探险者",
            normalize("86HIL"): "新蒙迪欧",
            normalize("86HJH"): "新蒙迪欧",
            normalize("87C3H"): "锐际",
        }
        return program, by_detail.get(detail_key, model)

    if program == "广汽传祺":
        by_detail = {
            normalize("GM6"): "GM6",
            normalize("GS4"): "GS4",
        }
        return program, by_detail.get(detail_key, model)

    if program == "华晨鑫源" and detail_key == normalize("JKC6420S6CXL5"):
        return program, "X30国六1"

    if program == "潍柴汽车" and detail_key == normalize("潍柴U70 2020款1.5T-6MT（7座）"):
        return program, "潍柴U70 （7座）"

    return program, model


def manual_dimension(program: str, common_name: str, model: str) -> Dimension | None:
    key = (normalize(program), normalize(common_name), normalize(model))
    program_key = normalize(program)
    model_key = normalize(model)
    if program_key == normalize("上汽通用五菱"):
        if model_key == normalize("LZW5022"):
            return Dimension(3797, 1820, "manual raw Sheet1 row 113")
        if model_key == normalize("LZW1020"):
            return Dimension(4358, 1811, "manual raw Sheet1 row 104")
        if model_key == normalize("LZW6449"):
            return Dimension(4420, 1770, "manual raw Sheet1 row 148")
    if program_key == normalize("东风小康"):
        if model_key in {normalize("B-风光500"), normalize("风光500")}:
            return Dimension(4385, 1650, "manual raw Sheet1 row 271")
        if model_key == normalize("B-风光580红星版"):
            return Dimension(4715, 1715, "manual raw Sheet1 row 270")
        if model_key == normalize("C-C52L"):
            return Dimension(5225, 1985, "manual raw Sheet1 row 272")
    if program_key == normalize("重庆长安"):
        if model_key in {normalize("SC7157ADB6"), normalize("SC7157ADD6.CNH4001.WC2")}:
            return Dimension(4900, 1500, "manual raw Sheet1 row 36")

    overrides: dict[tuple[str, str, str], Dimension] = {
        (normalize("上汽通用五菱"), normalize("五菱之光厢货"), normalize("LZW5022")): Dimension(3797, 1820, "manual raw Sheet1 row 113"),
        (normalize("上汽通用五菱"), normalize("五菱之光货车"), normalize("LZW1020")): Dimension(4358, 1811, "manual raw Sheet1 row 104"),
        (normalize("上汽通用五菱"), normalize("新款宏光S"), normalize("LZW6449")): Dimension(4420, 1770, "manual raw Sheet1 row 148"),
        (normalize("东风小康"), normalize("B-风光500"), normalize("B-风光500")): Dimension(4385, 1650, "manual raw Sheet1 row 271"),
        (normalize("东风小康"), normalize("风光500"), normalize("B-风光500")): Dimension(4385, 1650, "manual raw Sheet1 row 271"),
        (normalize("东风小康"), normalize("B-风光580红星版"), normalize("B-风光580红星版")): Dimension(4715, 1715, "manual raw Sheet1 row 270"),
        (normalize("东风小康"), normalize("C-C52L"), normalize("C-C52L")): Dimension(5225, 1985, "manual raw Sheet1 row 272"),
        (normalize("重庆长安"), normalize("锐骋CC蓝鲸版"), normalize("SC7157ADB6")): Dimension(4900, 1500, "manual raw Sheet1 row 36"),
        (normalize("重庆长安"), normalize("锐骋CC蓝鲸版"), normalize("SC7157ADD6.CNH4001.WC2")): Dimension(4900, 1500, "manual raw Sheet1 row 36"),
    }
    return overrides.get(key)


def load_raw_rows(path: Path) -> list[RawRow]:
    raw = pd.read_excel(path, sheet_name=0)
    raw["项目"] = raw["项目"].ffill()

    rows: list[RawRow] = []
    for idx, row in raw.iterrows():
        length = parse_single_int(row.get("长"))
        height = parse_single_int(row.get("高"))
        if length is None or height is None:
            continue
        rows.append(
            RawRow(
                row_number=idx + 2,
                project=str(row.get("项目", "")).strip(),
                common_name="" if pd.isna(row.get("俗称")) else str(row.get("俗称")).strip(),
                model="" if pd.isna(row.get("车型")) else str(row.get("车型")).strip(),
                length=length,
                height=height,
                norm_common=normalize(row.get("俗称")),
                norm_model=normalize(row.get("车型")),
                code_model=clean_code(row.get("车型")),
            )
        )
    return rows


def candidate_raw_rows(raw_rows: Iterable[RawRow], program: str, common_name: str, model: str) -> list[tuple[int, RawRow, str]]:
    norm_model = normalize(model)
    code_model = clean_code(model)
    norm_common = normalize(common_name)
    candidates: list[tuple[int, RawRow, str]] = []

    for raw in raw_rows:
        if not project_matches(raw.project, program):
            continue
        score = 0
        reasons: list[str] = []
        if raw.norm_common and raw.norm_common == norm_common:
            score = max(score, 100)
            reasons.append("common")
        if raw.norm_model and raw.norm_model == norm_model:
            score = max(score, 95)
            reasons.append("model-exact")
        if raw.norm_model and raw.norm_model == norm_common:
            score = max(score, 90)
            reasons.append("model-common")

        chunks = [chunk for chunk in re.split(r"[()/（）.,，;；、\\-]+", raw.norm_model) if chunk]
        if code_model and any(len(chunk) >= 3 and (chunk == code_model or chunk in norm_model) for chunk in chunks):
            score = max(score, 80)
            reasons.append("code-token")
        if norm_common and any(len(chunk) >= 2 and (chunk in norm_common or norm_common in chunk) for chunk in chunks):
            score = max(score, 70)
            reasons.append("common-token")

        if score:
            candidates.append((score, raw, "+".join(reasons)))

    if not candidates:
        return []
    best_score = max(score for score, _, _ in candidates)
    return [(score, raw, reason) for score, raw, reason in candidates if score == best_score]


def resolve_dimension(
    raw_rows: list[RawRow],
    fallback_dimensions: dict[tuple[str, str, str], tuple[int, int]],
    case_id: str,
    program: str,
    common_name: str,
    model: str,
) -> Dimension:
    manual = manual_dimension(program, common_name, model)
    if manual is not None:
        return manual

    candidates = candidate_raw_rows(raw_rows, program, common_name, model)
    unique_dims = sorted({(raw.length, raw.height) for _, raw, _ in candidates})
    if len(unique_dims) == 1:
        length, height = unique_dims[0]
        raw = candidates[0][1]
        reason = candidates[0][2]
        return Dimension(length, height, f"raw Sheet1 row {raw.row_number} {reason}")

    fallback = fallback_dimensions[(case_id, program, common_name)]
    if len(unique_dims) > 1:
        return Dimension(fallback[0], fallback[1], "kept current ambiguous raw dimensions")
    return Dimension(fallback[0], fallback[1], "kept current missing raw dimensions")


def load_case_csvs(real_case_dir: Path) -> dict[str, Path]:
    expected = {f"case_{idx:02d}" for idx in range(1, 11)}
    cases: dict[str, Path] = {}
    missing: list[str] = []

    for case_id in sorted(expected):
        candidates = [
            real_case_dir / case_id / f"{case_id}.csv",
            real_case_dir / f"{case_id}.csv",
            real_case_dir / f"{case_id}_cars.csv",
        ]
        existing = [path for path in candidates if path.exists()]
        if not existing:
            missing.append(case_id)
            continue
        cases[case_id] = existing[0]

    if missing:
        raise RuntimeError(f"Missing real-case CSV files for: {missing}")
    return cases


def load_original_case_frames(real_case_dir: Path) -> dict[str, pd.DataFrame]:
    archive_path = real_case_dir / "cars.zip"
    if not archive_path.exists():
        return {}

    frames: dict[str, pd.DataFrame] = {}
    with zipfile.ZipFile(archive_path) as archive:
        for name in archive.namelist():
            file_name = Path(name).name
            match = re.fullmatch(r"(case_\d{2})_cars\.csv", file_name)
            if match is None:
                continue
            with archive.open(name) as handle:
                frames[match.group(1)] = pd.read_csv(handle, encoding="utf-8-sig")
    return frames


def build_fallback_dimensions(
    case_files: dict[str, Path],
    vehicle_details: pd.DataFrame,
    original_case_frames: dict[str, pd.DataFrame],
) -> dict[tuple[str, str, str], tuple[int, int]]:
    fallback: dict[tuple[str, str, str], tuple[int, int]] = {}

    source_frames: dict[str, pd.DataFrame] = dict(original_case_frames)
    if not source_frames:
        source_frames = {
            case_id: pd.read_csv(path, encoding="utf-8-sig")
            for case_id, path in case_files.items()
        }

    for case_id, frame in source_frames.items():
        for _, row in frame.iterrows():
            fallback[(case_id, str(row["program"]), str(row["model"]))] = (
                int(row["length"]),
                int(row["height"]),
            )
    for _, row in vehicle_details.iterrows():
        program, model = restored_vehicle_type(str(row["厂商简称"]), str(row["俗称"]), str(row["型号"]))
        fallback.setdefault(
            (str(row["算例编号"]), program, model),
            (int(row["长"]), int(row["高"])),
        )
    return fallback


def load_case_order(
    case_files: dict[str, Path],
    original_case_frames: dict[str, pd.DataFrame],
) -> dict[str, dict[tuple[str, str], int]]:
    order: dict[str, dict[tuple[str, str], int]] = {}
    for case_id, path in case_files.items():
        frame = original_case_frames.get(case_id)
        if frame is None:
            frame = pd.read_csv(path, encoding="utf-8-sig")
        order[case_id] = {
            (str(row["program"]), str(row["model"])): idx
            for idx, row in frame.reset_index().iterrows()
        }
    return order


def corrected_vehicle_rows() -> tuple[pd.DataFrame, pd.DataFrame, dict[str, Path]]:
    case_files = load_case_csvs(REAL_CASE_DIR)
    original_case_frames = load_original_case_frames(REAL_CASE_DIR)
    details = pd.read_excel(SAMPLING_WORKBOOK, sheet_name="抽样车辆明细")
    fallback = build_fallback_dimensions(case_files, details, original_case_frames)
    raw_rows = load_raw_rows(RAW_DIMENSIONS)

    corrected_records: list[dict[str, object]] = []
    for row_index, row in details.iterrows():
        case_id = str(row["算例编号"])
        program = str(row["厂商简称"])
        common_name = str(row["俗称"])
        model = str(row["型号"])
        restored_program, restored_model = restored_vehicle_type(program, common_name, model)
        dimension = resolve_dimension(raw_rows, fallback, case_id, restored_program, restored_model, model)
        corrected_records.append(
            {
                "row_index": row_index,
                "case_id": case_id,
                "program": restored_program,
                "model": restored_model,
                "original_program": program,
                "original_model": common_name,
                "detail_model": model,
                "single_carriage_id": str(row["单次车厢ID"]),
                "layer": str(row["层号"]),
                "vin": str(row["车辆VIN号码"]),
                "width": int(row["宽"]),
                "length": dimension.length,
                "height": dimension.height,
                "old_length": int(row["长"]),
                "old_height": int(row["高"]),
                "source": dimension.source,
            }
        )

    vehicle_rows = pd.DataFrame(corrected_records)
    audit = vehicle_rows[
        [
            "case_id",
            "program",
            "model",
            "original_program",
            "original_model",
            "detail_model",
            "old_length",
            "old_height",
            "length",
            "height",
            "source",
        ]
    ].rename(columns={"length": "new_length", "height": "new_height"})
    audit["changed"] = (audit["old_length"] != audit["new_length"]) | (audit["old_height"] != audit["new_height"])

    return vehicle_rows, audit, case_files


def aggregate_case_cars(vehicle_rows: pd.DataFrame, case_order: dict[str, dict[tuple[str, str], int]]) -> dict[str, pd.DataFrame]:
    outputs: dict[str, pd.DataFrame] = {}
    for case_id, frame in vehicle_rows.groupby("case_id", sort=True):
        frame = frame.copy()
        frame["input_order"] = frame.apply(
            lambda row: case_order[case_id].get((row["program"], row["model"]), 10**6),
            axis=1,
        )
        grouped_rows: list[dict[str, object]] = []
        for (length, height), group in frame.sort_values(["input_order", "row_index"]).groupby(["length", "height"], sort=False):
            first = group.iloc[0]
            grouped_rows.append(
                {
                    "program": first["program"],
                    "model": first["model"],
                    "length": int(length),
                    "height": int(height),
                    "optional": 0,
                    "mandatory": int(len(group)),
                }
            )
        outputs[str(case_id)] = pd.DataFrame(grouped_rows, columns=["program", "model", "length", "height", "optional", "mandatory"])
    return outputs


def write_source_csvs(case_files: dict[str, Path], aggregated: dict[str, pd.DataFrame]) -> None:
    for case_id, path in case_files.items():
        path.parent.mkdir(parents=True, exist_ok=True)
        aggregated[case_id].to_csv(path, index=False, encoding="utf-8-sig")


def write_instance_dirs(aggregated: dict[str, pd.DataFrame], carriage_counts: dict[str, int]) -> None:
    for case_id, cars in aggregated.items():
        out_dir = REAL_CASE_DIR / case_id
        out_dir.mkdir(parents=True, exist_ok=True)
        instance_cars = cars.copy()
        instance_cars["optional"] = OPTIONAL_QUANTITY
        instance_cars.to_csv(out_dir / "cars.csv", index=False, encoding="utf-8-sig")
        pd.DataFrame([{"carriage_num": carriage_counts[case_id]}]).to_csv(
            out_dir / "carriage.csv",
            index=False,
            encoding="utf-8-sig",
        )


def load_carriage_counts() -> dict[str, int]:
    summary = pd.read_excel(SAMPLING_WORKBOOK, sheet_name="算例汇总")
    return {str(row["算例编号"]): int(row["抽取车厢数"]) for _, row in summary.iterrows()}


def update_workbook(vehicle_rows: pd.DataFrame) -> None:
    workbook = openpyxl.load_workbook(SAMPLING_WORKBOOK)
    rows_by_key = {
        (
            record.case_id,
            record.single_carriage_id,
            record.layer,
            record.vin,
        ): record
        for record in vehicle_rows.itertuples()
    }

    detail_sheet_names = ["抽样车辆明细"] + [f"case_{idx:02d}_车辆" for idx in range(1, 11)]
    for sheet_name in detail_sheet_names:
        sheet = workbook[sheet_name]
        headers = {sheet.cell(row=1, column=col).value: col for col in range(1, sheet.max_column + 1)}
        for excel_row in range(2, sheet.max_row + 1):
            key = (
                str(sheet.cell(excel_row, headers["算例编号"]).value),
                str(sheet.cell(excel_row, headers["单次车厢ID"]).value),
                str(sheet.cell(excel_row, headers["层号"]).value),
                str(sheet.cell(excel_row, headers["车辆VIN号码"]).value),
            )
            record = rows_by_key.get(key)
            if record is None:
                continue
            if "厂商简称" in headers:
                sheet.cell(excel_row, headers["厂商简称"]).value = str(record.program)
            if "俗称" in headers:
                sheet.cell(excel_row, headers["俗称"]).value = str(record.model)
            sheet.cell(excel_row, headers["长"]).value = int(record.length)
            sheet.cell(excel_row, headers["高"]).value = int(record.height)
            width = sheet.cell(excel_row, headers["宽"]).value
            if width is not None:
                sheet.cell(excel_row, headers["体积"]).value = int(record.length) * float(width) * int(record.height) / 1_000_000_000

    totals_by_case = vehicle_rows.groupby("case_id").agg(
        total_length=("length", "sum"),
        total_volume=("length", lambda _: 0.0),
    )
    type_count_by_case = vehicle_rows.drop_duplicates(["case_id", "length", "height"]).groupby("case_id").size()
    volume_by_case = (
        vehicle_rows.assign(volume=lambda df: df["length"] * df["width"] * df["height"] / 1_000_000_000)
        .groupby("case_id")["volume"]
        .sum()
    )
    totals_by_carriage = vehicle_rows.groupby(["case_id", "single_carriage_id"]).agg(total_length=("length", "sum"))
    volume_by_carriage = (
        vehicle_rows.assign(volume=lambda df: df["length"] * df["width"] * df["height"] / 1_000_000_000)
        .groupby(["case_id", "single_carriage_id"])["volume"]
        .sum()
    )
    upper_volume_by_carriage = (
        vehicle_rows[vehicle_rows["layer"] == "上层"]
        .assign(volume=lambda df: df["length"] * df["width"] * df["height"] / 1_000_000_000)
        .groupby(["case_id", "single_carriage_id"])["volume"]
        .sum()
    )
    type_summary_by_carriage: dict[tuple[str, str], tuple[int, str]] = {}
    for key, group in vehicle_rows.sort_values("row_index").groupby(["case_id", "single_carriage_id"], sort=False):
        unique_types: list[str] = []
        seen_dims: set[tuple[int, int]] = set()
        for _, record in group.iterrows():
            dim = (int(record["length"]), int(record["height"]))
            if dim in seen_dims:
                continue
            seen_dims.add(dim)
            unique_types.append(str(record["model"]))
        type_summary_by_carriage[(str(key[0]), str(key[1]))] = (len(seen_dims), "；".join(unique_types))

    summary = workbook["算例汇总"]
    summary_headers = {summary.cell(row=1, column=col).value: col for col in range(1, summary.max_column + 1)}
    for excel_row in range(2, summary.max_row + 1):
        case_id = str(summary.cell(excel_row, summary_headers["算例编号"]).value)
        if case_id not in totals_by_case.index:
            continue
        summary.cell(excel_row, summary_headers["总长度_mm"]).value = int(totals_by_case.loc[case_id, "total_length"])
        summary.cell(excel_row, summary_headers["总体积_m3"]).value = float(volume_by_case.loc[case_id])
        if "涉及车型数" in summary_headers:
            summary.cell(excel_row, summary_headers["涉及车型数"]).value = int(type_count_by_case.loc[case_id])

    carriage_sheet = workbook["抽样车厢清单"]
    carriage_headers = {carriage_sheet.cell(row=1, column=col).value: col for col in range(1, carriage_sheet.max_column + 1)}
    for excel_row in range(2, carriage_sheet.max_row + 1):
        case_id = str(carriage_sheet.cell(excel_row, carriage_headers["算例编号"]).value)
        single_carriage_id = str(carriage_sheet.cell(excel_row, carriage_headers["单次车厢ID"]).value)
        key = (case_id, single_carriage_id)
        if key not in totals_by_carriage.index:
            continue
        if "车型数" in carriage_headers and key in type_summary_by_carriage:
            carriage_sheet.cell(excel_row, carriage_headers["车型数"]).value = int(type_summary_by_carriage[key][0])
        if "车型清单" in carriage_headers and key in type_summary_by_carriage:
            carriage_sheet.cell(excel_row, carriage_headers["车型清单"]).value = type_summary_by_carriage[key][1]
        carriage_sheet.cell(excel_row, carriage_headers["总长度_mm"]).value = int(totals_by_carriage.loc[key, "total_length"])
        carriage_sheet.cell(excel_row, carriage_headers["总体积_m3"]).value = float(volume_by_carriage.loc[key])

    for sheet_name in detail_sheet_names:
        sheet = workbook[sheet_name]
        headers = {sheet.cell(row=1, column=col).value: col for col in range(1, sheet.max_column + 1)}
        for excel_row in range(2, sheet.max_row + 1):
            if sheet.cell(excel_row, headers["层号"]).value != "上层":
                sheet.cell(excel_row, headers["上层体积"]).value = None
                continue
            key = (
                str(sheet.cell(excel_row, headers["算例编号"]).value),
                str(sheet.cell(excel_row, headers["单次车厢ID"]).value),
            )
            if key in upper_volume_by_carriage.index:
                sheet.cell(excel_row, headers["上层体积"]).value = float(upper_volume_by_carriage.loc[key])

    workbook.save(SAMPLING_WORKBOOK)


def print_report(audit: pd.DataFrame, aggregated: dict[str, pd.DataFrame], carriage_counts: dict[str, int]) -> None:
    changed = audit[audit["changed"]].drop_duplicates(["case_id", "program", "model", "detail_model", "new_length", "new_height"])
    kept_missing = audit[audit["source"].str.contains("missing", regex=False)].drop_duplicates(["case_id", "program", "model", "detail_model"])
    kept_ambiguous = audit[audit["source"].str.contains("ambiguous", regex=False)].drop_duplicates(["case_id", "program", "model", "detail_model"])
    print(f"Corrected vehicle detail rows: {int(audit['changed'].sum())} / {len(audit)}")
    print(f"Changed detail groups: {len(changed)}")
    print(f"Kept missing raw groups: {len(kept_missing)}")
    print(f"Kept ambiguous raw groups: {len(kept_ambiguous)}")
    print()
    print("Generated cases:")
    for case_id in sorted(aggregated):
        cars = aggregated[case_id]
        print(
            f"  {case_id}: carriages={carriage_counts[case_id]}, "
            f"types={len(cars)}, mandatory={int(cars['mandatory'].sum())}, "
            f"length_sum={int((cars['length'] * cars['mandatory']).sum())}"
        )


def validate_outputs(aggregated: dict[str, pd.DataFrame], carriage_counts: dict[str, int]) -> None:
    for case_id, cars in aggregated.items():
        if not (cars["mandatory"] > 0).all():
            raise RuntimeError(f"{case_id}: mandatory quantities must be positive after aggregation")
        if cars[["length", "height"]].duplicated().any():
            raise RuntimeError(f"{case_id}: duplicate dimensions remain after aggregation")
        instance_cars_path = REAL_CASE_DIR / case_id / "cars.csv"
        instance_carriage_path = REAL_CASE_DIR / case_id / "carriage.csv"
        if instance_cars_path.exists():
            instance = pd.read_csv(instance_cars_path, encoding="utf-8-sig")
            if not (instance["optional"] == OPTIONAL_QUANTITY).all():
                raise RuntimeError(f"{case_id}: generated optional quantities are not fixed to {OPTIONAL_QUANTITY}")
        if instance_carriage_path.exists():
            carriage = pd.read_csv(instance_carriage_path, encoding="utf-8-sig")
            if int(carriage.loc[0, "carriage_num"]) != carriage_counts[case_id]:
                raise RuntimeError(f"{case_id}: carriage count mismatch")


def main() -> None:
    parser = argparse.ArgumentParser(description="Prepare corrected real-case motorail instances.")
    parser.add_argument("--dry-run", action="store_true", help="Only report changes; do not write files.")
    args = parser.parse_args()

    vehicle_rows, audit, case_files = corrected_vehicle_rows()
    original_case_frames = load_original_case_frames(REAL_CASE_DIR)
    case_order = load_case_order(case_files, original_case_frames)
    aggregated = aggregate_case_cars(vehicle_rows, case_order)
    carriage_counts = load_carriage_counts()

    print_report(audit, aggregated, carriage_counts)

    if args.dry_run:
        return

    write_source_csvs(case_files, aggregated)
    write_instance_dirs(aggregated, carriage_counts)
    update_workbook(vehicle_rows)
    validate_outputs(aggregated, carriage_counts)


if __name__ == "__main__":
    main()
